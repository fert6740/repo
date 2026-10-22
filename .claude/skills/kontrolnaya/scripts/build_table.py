"""Собрать итоговые ответы всех чатов в Excel-таблицу.

Использование:
    python3 -I build_table.py <ответы.json> <выход.xlsx>

Формат JSON:
{
  "subject": "Литература",
  "rows": [
    {"n": 1, "task": "кратко, о чём задание", "answer": "итоговый ответ", "confidence": 0.99,
     "level": "высокая", "why": "что снижает уверенность", "session_id": "session_..."}
  ]
}
confidence — доля (0.85 = 85 %). Строки без ответа (чат упал/не закончил): answer = "НЕТ ОТВЕТА",
confidence = null — такие строки подсвечиваются красным.
"""
import glob
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
SESSION_URL = "https://claude.ai/code/"


def main(src, out):
    data = json.load(open(src, encoding="utf-8"))
    rows = sorted(data["rows"], key=lambda r: r["n"])

    wb = Workbook()
    ws = wb.active
    ws.title = data["subject"][:31]
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(["№", "Задание (кратко)", "Итоговый ответ", "Уверенность", "Уровень",
               "Что снижает уверенность", "Чат"])
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    for r in rows:
        ws.append([r["n"], r.get("task", ""), r.get("answer", ""), r.get("confidence"),
                   r.get("level", ""), r.get("why", ""), f"Открыть чат {r['n']}"])
        i = ws.max_row
        ws.cell(i, 7).hyperlink = SESSION_URL + r["session_id"]
        for col in range(1, 8):
            c = ws.cell(i, col)
            c.border = border
            c.font = (Font(name=FONT, color="0563C1", underline="single") if col == 7
                      else Font(name=FONT, bold=(col == 3)))
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if col in (1, 4, 5) else "left")
        ws.cell(i, 4).number_format = "0%"
        if r.get("confidence") is None:
            for col in range(1, 7):
                ws.cell(i, col).fill = PatternFill("solid", fgColor="F8CBAD")

    last = ws.max_row
    ws.append(["", "Средняя уверенность", "", f"=AVERAGE(D2:D{last})", "", "", ""])
    i = ws.max_row
    for col in range(1, 8):
        ws.cell(i, col).font = Font(name=FONT, bold=True)
        ws.cell(i, col).border = border
    ws.cell(i, 4).number_format = "0%"
    ws.cell(i, 4).alignment = Alignment(horizontal="center")

    # уверенность ниже 90 % — жёлтым
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(
        operator="lessThan", formula=["0.9"], fill=PatternFill("solid", fgColor="FFF2CC")))
    for col, w in zip("ABCDEFG", [5, 38, 60, 13, 11, 50, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(out)

    # пересчитать формулу средней уверенности, если доступен recalc.py из скилла xlsx
    recalc = glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True)
    if recalc:
        skill_dir = recalc[0].rsplit("/scripts/", 1)[0]
        print(subprocess.run([sys.executable, recalc[0], os.path.abspath(out)],
                             capture_output=True, text=True, cwd=skill_dir).stdout)
    else:
        print("recalc.py не найден: среднее посчитается при открытии в Excel")
    print(f"Готово: {out}, строк: {len(rows)}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
